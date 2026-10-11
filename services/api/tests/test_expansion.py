"""Acceptance tests for v0.2: inspect actual results, not just HTTP status codes."""
import io
import json
import time
import zipfile
from dataclasses import replace
from pathlib import Path
import pytest
from fastapi.testclient import TestClient
from PIL import Image
from pypdf import PdfReader, PdfWriter
from pypdf.generic import RectangleObject, NameObject, ArrayObject, DictionaryObject
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from docx import Document
from app.access import hash_password, verify_password, WindowLimit
from app.config import Settings
from app.engines import convert
from app.extra_engines import safe_text
from app.jobs import JobStore, QueueFull
from app.main import create_app, display_name
from app.models import ConvertOptions, ConversionError
from app.pipeline import execute
from test_api import HEADERS, wait_for_job


def converted(tool, file, tmp_path, settings, **options):
    return convert(tool, [file], tmp_path / tool, ConvertOptions(**options), settings, lambda *_: None)


def test_organizer_preserves_text_order_and_rotation(sample_pdf, tmp_path, settings):
    output, _ = converted('organize-pdf', sample_pdf, tmp_path, settings,
        page_plan=[{'page': 3, 'rotation': 90}, {'page': 1, 'rotation': 270}])
    pages = PdfReader(output).pages
    assert len(pages) == 2
    assert 'page 3' in pages[0].extract_text() and pages[0].rotation == 90
    assert 'page 1' in pages[1].extract_text() and pages[1].rotation == 270
    for plan in ([], [{'page': 4}], [{'page': 1}, {'page': 1}]):
        with pytest.raises(ConversionError):
            converted('organize-pdf', sample_pdf, tmp_path, settings, page_plan=plan)


def test_crop_is_visual_not_redaction(sample_pdf, tmp_path, settings):
    output, warnings = converted('crop-pdf', sample_pdf, tmp_path, settings, pages='2',
        crop_left=10, crop_right=20, crop_top=5, crop_bottom=15)
    pages = PdfReader(output).pages
    assert list(pages[0].cropbox) == [0, 0, 720, 405]
    assert list(map(float, pages[1].cropbox)) == pytest.approx([72, 60.75, 576, 384.75])
    assert 'page 2' in pages[1].extract_text() and any('восстановлено' in w for w in warnings)


def test_text_exports_are_honest_and_ordered(sample_pdf, tmp_path, settings):
    output, warnings = converted('pdf-to-docx', sample_pdf, tmp_path, settings, pages='3,1')
    text = '\n'.join(p.text for p in Document(output).paragraphs)
    assert text.index('page 3') < text.index('page 1')
    assert 'page 2' not in text and any('не восстановлены' in w for w in warnings)
    md, _ = converted('pdf-to-markdown', sample_pdf, tmp_path, settings, pages='2')
    assert '## Страница 2' in md.read_text() and 'page 2' in md.read_text()
    assert safe_text('a\x00\x0b\ud800b') == 'ab'


def test_extract_embedded_images_and_rasterize(tmp_path, settings):
    source = tmp_path / 'embedded.pdf'
    pdf = canvas.Canvas(str(source), pagesize=(360, 240))
    image = Image.new('RGB', (80, 40), (20, 80, 150))
    pdf.drawImage(ImageReader(image), 30, 50, 160, 80)
    pdf.drawString(30, 180, 'Text must not survive rasterization')
    pdf.save(); image.close()
    extracted, _ = converted('pdf-extract-images', source, tmp_path, settings)
    with zipfile.ZipFile(extracted) as archive:
        assert len(archive.namelist()) == 1
        with Image.open(io.BytesIO(archive.read(archive.namelist()[0]))) as result:
            assert result.size == (80, 40)
    flattened, warnings = converted('flatten-pdf', source, tmp_path, settings, dpi=96)
    page = PdfReader(flattened).pages[0]
    assert not page.extract_text().strip() and len(page.images) == 1
    assert float(page.mediabox.width) == 360 and warnings


def test_diagnostics_resolves_indirect_annotations(sample_pdf, tmp_path, settings):
    writer = PdfWriter(clone_from=sample_pdf)
    writer.pages[0][NameObject('/Annots')] = writer._add_object(ArrayObject())
    indirect = tmp_path / 'indirect.pdf'; writer.write(indirect)
    report, warnings = converted('pdf-info', indirect, tmp_path, settings)
    data = json.loads(report.read_text())
    assert data['pageCount'] == 3 and data['pages'][0]['annotations'] == 0 and warnings


@pytest.mark.parametrize('tool', ['image-rotate', 'image-flip', 'image-grayscale'])
def test_image_operations_preserve_alpha(tool, tmp_path, settings):
    image = Image.new('RGBA', (8, 4), (200, 80, 20, 100))
    image.putpixel((0, 0), (10, 20, 30, 230))
    source = tmp_path / 'pixels.png'; image.save(source); image.close()
    result, _ = converted(tool, source, tmp_path, settings, format='png', angle=90)
    with Image.open(result) as output:
        assert output.mode == 'RGBA'
        if tool == 'image-rotate':
            assert output.size == (4, 8) and output.getpixel((3, 0)) == (10, 20, 30, 230)
        elif tool == 'image-flip':
            assert output.getpixel((7, 0)) == (10, 20, 30, 230)
        else:
            r, g, b, a = output.getpixel((0, 0))
            assert r == g == b and a == 230


def test_split_groups_and_selected_presentation(sample_pdf, tmp_path, settings):
    output, _ = converted('split-pdf', sample_pdf, tmp_path, settings, pages='3,1,2', split_every=2)
    with zipfile.ZipFile(output) as archive:
        assert len(archive.namelist()) == 2
        group = PdfReader(io.BytesIO(archive.read(archive.namelist()[0])))
        assert len(group.pages) == 2 and 'page 3' in group.pages[0].extract_text()
    from pptx import Presentation
    pptx, _ = converted('pdf-to-pptx', sample_pdf, tmp_path, settings, pages='2', dpi=96)
    assert len(Presentation(pptx).slides) == 1


def test_stamps_on_offset_rotated_and_selected_pages(sample_pdf, tmp_path, settings):
    writer = PdfWriter(clone_from=sample_pdf)
    writer.pages[0].mediabox = RectangleObject((-100, -50, 620, 355))
    writer.pages[0].cropbox = RectangleObject((-80, -30, 600, 335))
    writer.pages[1].rotate(90)
    source = tmp_path / 'offset.pdf'; writer.write(source)
    output, _ = converted('watermark-pdf', source, tmp_path, settings, text='REVIEW', pages='1-2', watermark_angle=0)
    pages = PdfReader(output).pages
    assert all('REVIEW' in pages[i].extract_text() for i in (0, 1))
    assert 'REVIEW' not in pages[2].extract_text()
    import pypdfium2 as pdfium
    with pdfium.PdfDocument(output) as document:
        for index in range(3):
            page = document[index]
            try:
                bitmap = page.render(scale=1)
                assert bitmap.width > 0 and bitmap.height > 0
                bitmap.close()
            finally:
                page.close()


def test_batch_partial_has_machine_readable_report(sample_pdf, tmp_path, settings):
    bad = tmp_path / 'broken.pdf'; bad.write_bytes(b'not pdf')
    output, warnings, summary = execute('rotate-pdf', [sample_pdf, bad], ['../safe.pdf', 'broken.pdf'],
        tmp_path / 'batch', ConvertOptions(batch=True), settings, lambda *_: None)
    assert summary == {'processed': 1, 'total': 2, 'failed': 1}
    assert any('Часть' in warning for warning in warnings)
    with zipfile.ZipFile(output) as archive:
        assert all('/' not in name and '\\' not in name for name in archive.namelist())
        report = json.loads(archive.read('report.json'))
        assert [item['status'] for item in report['files']] == ['done', 'error']
        assert PdfReader(io.BytesIO(archive.read(report['files'][0]['result']))).pages[0].rotation == 90
    assert not (output.parent / '.work').exists()


def test_postprocess_is_applied_to_real_output(sample_image, tmp_path, settings):
    output, _, batch = execute('images-to-pdf', [sample_image], ['image.png'], tmp_path / 'finished',
        ConvertOptions(postprocess='watermark', text='APPROVED'), settings, lambda *_: None)
    assert batch is None and 'APPROVED' in PdfReader(output).pages[0].extract_text()


def test_continuation_ownership_and_history_clear(sample_pdf, settings):
    app = create_app(settings)
    with TestClient(app) as client:
        response = client.post('/api/v1/jobs', headers=HEADERS, data={'tool': 'rotate-pdf'}, files={'files': ('a.pdf', sample_pdf.read_bytes())})
        source = wait_for_job(client, response.json()['id'])
        endpoint = f"/api/v1/jobs/{source['id']}/continue"
        other = TestClient(app)
        assert other.post(endpoint, headers=HEADERS, json={'tool': 'pdf-to-text'}).status_code == 404
        created = client.post(endpoint, headers=HEADERS, json={'tool': 'pdf-to-text'})
        assert created.status_code == 202, created.text
        result = wait_for_job(client, created.json()['id'])
        assert result['state'] == 'done' and 'page 1' in client.get(result['downloadUrl']).text
        assert client.post(endpoint, headers=HEADERS, json={'tool': 'image-convert'}).status_code == 422
        assert other.post('/api/v1/history/clear', headers=HEADERS).json()['deleted'] == 0
        assert client.post('/api/v1/history/clear', headers=HEADERS).json()['deleted'] == 2
        assert client.get(source['downloadUrl']).status_code == 404


def test_batch_summary_exposed_without_passwords(sample_pdf, settings):
    with TestClient(create_app(settings)) as client:
        response = client.post('/api/v1/jobs', headers=HEADERS,
            data={'tool': 'rotate-pdf', 'options': json.dumps({'batch': True})},
            files=[('files', ('good.pdf', sample_pdf.read_bytes())), ('files', ('bad.pdf', b'bad'))])
        job = wait_for_job(client, response.json()['id'])
        assert job['state'] == 'done' and job['batch']['failed'] == 1
        assert job['stage'] == 'Готово с ошибками' and 'options' not in job


def test_password_gate_protects_operations_and_schema(sample_pdf, settings):
    password = 'test-password-1234'
    encoded = hash_password(password)
    assert verify_password(password, encoded) and not verify_password('wrong', encoded)
    assert not verify_password('\ud800', encoded)
    app = create_app(replace(settings, access_password_hash=encoded))
    with TestClient(app) as client:
        assert client.get('/api/health').status_code == 200
        assert client.get('/api/v1/access').json() == {'enabled': True, 'authenticated': False}
        for route in ('/api/v1/config', '/api/v1/jobs', '/api/openapi.json'):
            assert client.get(route).status_code == 401
        assert client.post('/api/v1/access/login', headers=HEADERS, json={'password': 'wrong'}).status_code == 401
        assert client.post('/api/v1/access/login', headers=HEADERS, json={'password': password}).status_code == 200
        assert client.get('/api/v1/config').json()['accessEnabled'] is True
        assert client.get('/api/openapi.json').status_code == 200
        assert client.post('/api/v1/access/logout', headers=HEADERS).status_code == 200
        assert client.get('/api/v1/jobs').status_code == 401


def test_request_hardening_and_large_page_plan(settings, sample_pdf):
    app = create_app(replace(settings, max_pages=1000))
    client = TestClient(app)  # Deliberately leave the queue stopped: this tests admission only.
    assert client.get('/api/health').status_code == 503
    for origin in ('null', 'http://[', 'https://testserver/path', 'https://testserver#bad'):
        assert client.post('/api/v1/history/clear', headers={**HEADERS, 'Origin': origin}).status_code == 403
    plan = [{'page': i, 'rotation': 0} for i in range(1, 1001)]
    response = client.post('/api/v1/jobs', headers=HEADERS,
        data={'tool': 'organize-pdf', 'options': json.dumps({'page_plan': plan})},
        files={'files': ('large-plan.pdf', sample_pdf.read_bytes())})
    assert response.status_code == 202, response.text
    for options in ({'angle': 12}, {'postprocess': 'watermark'}, {'format': 'png'}):
        response = client.post('/api/v1/jobs', headers=HEADERS,
            data={'tool': 'rotate-pdf', 'options': json.dumps(options)}, files={'files': ('a.pdf', sample_pdf.read_bytes())})
        assert response.status_code == 422
    too_many = [('files', (f'{i}.pdf', b'fake')) for i in range(21)]
    assert client.post('/api/v1/jobs', headers=HEADERS, data={'tool': 'merge-pdf'}, files=too_many).status_code == 400
    assert display_name('a' * 220 + '.pdf').endswith('.pdf')
    assert '\u202e' not in display_name('report\u202e.pdf')


def test_storage_reservation_recovery_and_limiter(settings):
    store = JobStore(replace(settings, session_storage_mb=256))
    jid = 'a' * 32
    store.create(jid, 'session', 'rotate-pdf', ['a.pdf'], 10)
    with pytest.raises(QueueFull):
        store.create('b' * 32, 'session', 'rotate-pdf', ['b.pdf'], 10)
    store.claim()
    root = store.root / jid
    for directory in ('input', 'output', 'scratch'):
        (root / directory).mkdir(parents=True)
        (root / directory / 'private').write_text('synthetic')
    (root / 'manifest.json').write_text('{}')
    store.recover()
    assert store.get(jid)['state'] == 'error'
    assert not any((root / directory).exists() for directory in ('input', 'output', 'scratch'))
    assert not (root / 'manifest.json').exists()
    limit = WindowLimit(2, 60, max_keys=2)
    assert limit.take('one') and limit.take('one') and not limit.take('one')
    limit.take('two'); limit.take('three')
    assert len(limit.events) == 2


def render_rgb(path, page_index, dpi=96):
    import pypdfium2 as pdfium
    document = pdfium.PdfDocument(path)
    page = document[page_index]
    bitmap = None
    try:
        bitmap = page.render(scale=dpi / 72, fill_color=(255, 255, 255, 255))
        return bitmap.to_pil().convert('RGB').copy()
    finally:
        if bitmap is not None:
            bitmap.close()
        page.close(); document.close()


def test_reordering_does_not_change_rendered_page_pixels(sample_pdf, tmp_path, settings):
    from PIL import ImageChops
    output, _ = converted('organize-pdf', sample_pdf, tmp_path, settings,
        page_plan=[{'page': 3}, {'page': 1}, {'page': 2}])
    for position, original in enumerate([2, 0, 1]):
        expected = render_rgb(sample_pdf, original)
        actual = render_rgb(output, position)
        try:
            assert expected.size == actual.size
            assert ImageChops.difference(expected, actual).getbbox() is None
        finally:
            expected.close(); actual.close()


def test_presentation_images_match_source_pdf_renders(sample_pdf, tmp_path, settings):
    from PIL import ImageChops
    from pptx import Presentation
    output, _ = converted('pdf-to-pptx', sample_pdf, tmp_path, settings, dpi=96)
    for index, slide in enumerate(Presentation(output).slides):
        expected = render_rgb(sample_pdf, index)
        with Image.open(io.BytesIO(slide.shapes[0].image.blob)) as embedded:
            actual = embedded.convert('RGB')
            try:
                assert expected.size == actual.size
                assert ImageChops.difference(expected, actual).getbbox() is None
            finally:
                expected.close(); actual.close()
