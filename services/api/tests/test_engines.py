import io
import shutil
import zipfile
import pytest
from PIL import Image
from pypdf import PdfReader
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches
from app.engines import convert, validate_container
from app.models import ConversionError, ConvertOptions
from app.pages import parse_pages

@pytest.mark.parametrize(('value','expected'),[('',[0,1,2]),('3,1-2',[2,0,1]),('1–2',[0,1]),('2',[1])])
def test_ranges(value,expected):
    assert parse_pages(value,3)==expected

@pytest.mark.parametrize('value',['0','4','2-1','1,1','a','1-','1,,2','-1','1;2'])
def test_bad_ranges(value):
    with pytest.raises(ConversionError): parse_pages(value,3)

def run(tool,files,tmp_path,settings,**opts):
    events=[]
    output,warnings=convert(tool,files,tmp_path/tool,ConvertOptions(**opts),settings,lambda value,stage:events.append((value,stage)))
    assert output.is_file() and output.stat().st_size>0 and events
    return output,warnings

def test_pdf_operations(sample_pdf,tmp_path,settings):
    merged,_=run('merge-pdf',[sample_pdf,sample_pdf],tmp_path,settings)
    assert len(PdfReader(merged).pages)==6
    split,_=run('split-pdf',[sample_pdf],tmp_path,settings)
    with zipfile.ZipFile(split) as a:
        assert len(a.namelist())==3
        assert len(PdfReader(io.BytesIO(a.read('page-001.pdf'))).pages)==1
    extracted,_=run('extract-pages',[sample_pdf],tmp_path,settings,pages='3,1')
    assert 'page 3' in PdfReader(extracted).pages[0].extract_text()
    removed,_=run('remove-pages',[sample_pdf],tmp_path,settings,pages='2')
    assert len(PdfReader(removed).pages)==2
    rotated,_=run('rotate-pdf',[sample_pdf],tmp_path,settings,angle=90,pages='2')
    reader=PdfReader(rotated)
    assert reader.pages[0].rotation==0 and reader.pages[1].rotation==90
    compressed,_=run('compress-pdf',[sample_pdf],tmp_path,settings)
    assert compressed.stat().st_size<=sample_pdf.stat().st_size
    assert len(PdfReader(compressed).pages)==3
    text,_=run('pdf-to-text',[sample_pdf],tmp_path,settings)
    assert 'page 1' in text.read_text()

def test_passwords(sample_pdf,tmp_path,settings):
    protected,_=run('protect-pdf',[sample_pdf],tmp_path,settings,new_password='test-pass-123')
    reader=PdfReader(protected); assert reader.is_encrypted and reader.decrypt('test-pass-123')
    unlocked,_=run('unlock-pdf',[protected],tmp_path,settings,password='test-pass-123')
    assert not PdfReader(unlocked).is_encrypted
    with pytest.raises(ConversionError): run('unlock-pdf',[protected],tmp_path,settings,password='wrong')
    with pytest.raises(ConversionError): run('protect-pdf',[sample_pdf],tmp_path,settings,new_password='short')

def test_stamps_and_metadata(sample_pdf,tmp_path,settings):
    numbered,_=run('number-pages',[sample_pdf],tmp_path,settings,start=10)
    assert '10' in PdfReader(numbered).pages[0].extract_text()
    marked,_=run('watermark-pdf',[sample_pdf],tmp_path,settings,text='IMONSOFT TEST')
    assert 'IMONSOFT TEST' in PdfReader(marked).pages[0].extract_text()
    clean,_=run('pdf-metadata',[sample_pdf],tmp_path,settings)
    assert not PdfReader(clean).metadata

def test_pdf_to_pptx_is_image_based(sample_pdf,tmp_path,settings):
    path,warnings=run('pdf-to-pptx',[sample_pdf],tmp_path,settings,dpi=96)
    deck=Presentation(path)
    assert len(deck.slides)==3 and warnings
    assert deck.slide_width/deck.slide_height==pytest.approx(16/9)
    for slide in deck.slides:
        assert len(slide.shapes)==1
        assert slide.shapes[0].shape_type==MSO_SHAPE_TYPE.PICTURE
        assert slide.shapes[0].left==0 and slide.shapes[0].top==0

@pytest.mark.parametrize('tool',['pdf-to-jpg','pdf-to-png'])
def test_page_images(tool,sample_pdf,tmp_path,settings):
    output,_=run(tool,[sample_pdf],tmp_path,settings,dpi=96)
    with zipfile.ZipFile(output) as a:
        assert len(a.namelist())==3
        with Image.open(io.BytesIO(a.read(a.namelist()[0]))) as im: assert im.size==(960,540)

def test_image_workflows(sample_image,tmp_path,settings):
    path,_=run('images-to-pdf',[sample_image,sample_image],tmp_path,settings)
    assert len(PdfReader(path).pages)==2
    for fmt in ('jpg','png','webp'):
        path,_=run('image-convert',[sample_image],tmp_path,settings,format=fmt)
        with Image.open(path) as im: assert im.size==(320,200)
    resized,_=run('image-resize',[sample_image],tmp_path,settings,width=100,height=100)
    with Image.open(resized) as im: assert im.size in ((100,63),(100,62))
    path,_=run('image-compress',[sample_image],tmp_path,settings,format='webp')
    assert path.suffix=='.webp'

def test_invalid_files(sample_pdf,tmp_path,settings):
    wrong=tmp_path/'fake.pdf';wrong.write_text('not a pdf')
    with pytest.raises(ConversionError): validate_container(wrong)
    with pytest.raises(ConversionError): run('remove-pages',[sample_pdf],tmp_path,settings,pages='1-3')
    wrong=tmp_path/'fake.pptx'
    with zipfile.ZipFile(wrong,'w') as a:a.writestr('word/document.xml','<document/>')
    with pytest.raises(ConversionError):validate_container(wrong)

@pytest.mark.skipif(not shutil.which('libreoffice'),reason='LibreOffice not installed')
def test_real_powerpoint_to_pdf(tmp_path,settings):
    deck=Presentation();deck.slide_width,deck.slide_height=Inches(16),Inches(9)
    slide=deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1),Inches(1),Inches(12),Inches(2)).text='IMONSOFT conversion acceptance'
    source=tmp_path/'source.pptx';deck.save(source)
    output,_=run('pptx-to-pdf',[source],tmp_path,settings)
    reader=PdfReader(output)
    assert len(reader.pages)==1 and 'IMONSOFT' in reader.pages[0].extract_text()
    assert float(reader.pages[0].mediabox.width)/float(reader.pages[0].mediabox.height)==pytest.approx(16/9,rel=.001)

@pytest.mark.skipif(not shutil.which('tesseract'),reason='Tesseract not installed')
def test_ocr(sample_pdf,tmp_path,settings):
    path,warnings=run('ocr-pdf',[sample_pdf],tmp_path,settings,language='eng')
    reader=PdfReader(path)
    assert len(reader.pages)==3 and warnings
    assert 'IMONSOFT' in reader.pages[0].extract_text()
