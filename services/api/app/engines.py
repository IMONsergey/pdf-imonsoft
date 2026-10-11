from __future__ import annotations

import io
import math
import os
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from typing import Callable

import httpx
from PIL import Image, ImageOps, UnidentifiedImageError
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader

from .config import Settings
from .models import ConversionError, ConvertOptions
from .pages import parse_pages

Progress = Callable[[int | None, str], None]
Image.MAX_IMAGE_PIXELS = 40_000_000


def validate_container(path: Path) -> None:
    """Reject malformed files and oversized Office ZIPs before calling native parsers."""
    suffix = path.suffix.lower()
    with path.open('rb') as stream:
        header = stream.read(1024)
    if suffix == '.pdf':
        if b'%PDF-' not in header:
            raise ConversionError('Файл не является PDF. Изменение расширения не конвертирует документ.')
    elif suffix in ('.pptx', '.docx', '.xlsx', '.odt', '.ods', '.odp'):
        try:
            with zipfile.ZipFile(path) as archive:
                entries = archive.infolist()
                if len(entries) > 10000 or sum(e.file_size for e in entries) > 250 * 1024 * 1024:
                    raise ConversionError('Распакованный документ слишком большой.')
                for entry in entries:
                    if entry.file_size > 1024 * 1024 and entry.file_size / max(1, entry.compress_size) > 500:
                        raise ConversionError('Внутренняя структура файла превышает безопасные лимиты.')
                    if entry.flag_bits & 1:
                        raise ConversionError('Зашифрованный Office-архив не поддерживается.')
                required = {'.pptx': 'ppt/presentation.xml', '.docx': 'word/document.xml', '.xlsx': 'xl/workbook.xml'}
                if suffix in required and required[suffix] not in archive.namelist():
                    raise ConversionError('Содержимое файла не соответствует его расширению.')
                if suffix in ('.odt', '.ods', '.odp'):
                    expected = {'.odt': b'application/vnd.oasis.opendocument.text', '.ods': b'application/vnd.oasis.opendocument.spreadsheet', '.odp': b'application/vnd.oasis.opendocument.presentation'}
                    if archive.read('mimetype').strip() != expected[suffix]:
                        raise ConversionError('Содержимое OpenDocument не соответствует его расширению.')
        except (zipfile.BadZipFile, KeyError):
            raise ConversionError('Не удалось открыть структуру Office-документа.') from None
    elif suffix in ('.ppt', '.doc', '.xls'):
        if not header.startswith(bytes.fromhex('D0CF11E0A1B11AE1')):
            raise ConversionError('Неверный формат старого Office-документа.')
    elif suffix == '.rtf' and not header.lstrip().startswith(b'{\\rtf'):
        raise ConversionError('Файл не является документом RTF.')


def reader_for(path: Path, options: ConvertOptions, settings: Settings) -> PdfReader:
    try:
        reader = PdfReader(path)
        if reader.is_encrypted and not reader.decrypt(options.password):
            raise ConversionError('Необходим верный пароль для открытия PDF.')
        count = len(reader.pages)
    except ConversionError:
        raise
    except Exception:
        raise ConversionError('Не удалось прочитать PDF. Проверьте файл и пароль.') from None
    if count == 0 or count > settings.max_pages:
        raise ConversionError(f'Допустимо от 1 до {settings.max_pages} страниц в документе.')
    return reader


def write_pdf(writer: PdfWriter, output: Path) -> Path:
    with output.open('wb') as handle:
        writer.write(handle)
    writer.close()
    return output


def make_zip(files: list[Path], output: Path) -> Path:
    with zipfile.ZipFile(output, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        for file in files:
            archive.write(file, arcname=file.name)
    for file in files:
        file.unlink(missing_ok=True)
    return output


def stamp_font() -> str:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    candidates = [os.getenv('STAMP_FONT_PATH', ''), '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', '/Library/Fonts/Arial Unicode.ttf', '/System/Library/Fonts/Supplemental/Arial.ttf']
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            if 'IMONStamp' not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont('IMONStamp', candidate))
            return 'IMONStamp'
    return 'Helvetica'


def render_pages(path: Path, indexes: list[int], dpi: int, password: str, progress: Progress):
    # PDFium is not thread-safe. This code runs in one dedicated process per job.
    import pypdfium2 as pdfium
    document = pdfium.PdfDocument(path, password=password or None)
    total_pixels = 0
    try:
        for position, index in enumerate(indexes):
            page = document[index]
            bitmap = None
            try:
                width, height = page.get_size()
                pixels = math.ceil(width * dpi / 72) * math.ceil(height * dpi / 72)
                total_pixels += pixels
                if width <= 0 or height <= 0 or pixels > 24_000_000 or total_pixels > 300_000_000:
                    raise ConversionError('Слишком большое разрешение или число страниц. Уменьшите DPI или выберите часть документа.')
                bitmap = page.render(scale=dpi / 72, fill_color=(255, 255, 255, 255))
                # Detach the image before closing native handles.
                image = bitmap.to_pil().convert('RGB').copy()
            finally:
                if bitmap is not None:
                    bitmap.close()
                page.close()
            progress(10 + round(75 * (position + 1) / len(indexes)), f'Страница {position + 1} из {len(indexes)}')
            yield image, (width, height)
    finally:
        document.close()


def office_to_pdf(path: Path, outdir: Path, settings: Settings, progress: Progress) -> Path:
    output = outdir / 'result.pdf'
    progress(None, 'Конвертация в LibreOffice')
    if settings.office_engine == 'local':
        # Development-only adapter for trusted fixtures. Production uses an isolated Gotenberg.
        executable = shutil.which(settings.libreoffice_bin) or settings.libreoffice_bin
        with tempfile.TemporaryDirectory(dir=outdir, prefix='office-profile-') as profile:
            try:
                process = subprocess.run(
                    [executable, f'-env:UserInstallation={Path(profile).as_uri()}', '--headless', '--convert-to', 'pdf', '--outdir', str(outdir), str(path)],
                    stdin=subprocess.DEVNULL, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                    timeout=max(10, settings.job_timeout - 10), check=False,
                )
            except (OSError, subprocess.TimeoutExpired):
                raise ConversionError('Локальный LibreOffice недоступен или не завершил конвертацию.') from None
            source = outdir / f'{path.stem}.pdf'
            if process.returncode != 0 or not source.is_file():
                raise ConversionError('LibreOffice не смог преобразовать этот документ.')
            source.replace(output)
    else:
        try:
            with path.open('rb') as stream, httpx.Client(timeout=settings.job_timeout - 5, trust_env=False) as client:
                with client.stream('POST', settings.gotenberg_url + '/forms/libreoffice/convert', files={'files': (path.name, stream)}, data={'exportFormFields': 'false'}) as response:
                    if response.status_code != 200:
                        raise ConversionError('Office-движок отклонил документ. Проверьте его формат и содержимое.')
                    total = 0
                    with output.open('wb') as target:
                        for chunk in response.iter_bytes(1024 * 1024):
                            total += len(chunk)
                            if total > settings.max_output_mb * 1024 * 1024:
                                raise ConversionError('Результат превышает допустимый размер.')
                            target.write(chunk)
        except httpx.HTTPError:
            raise ConversionError('Office-движок недоступен или время обработки истекло. Проверьте Gotenberg.') from None
    if not output.is_file() or not output.read_bytes()[:5] == b'%PDF-':
        raise ConversionError('Office-движок не вернул корректный PDF.')
    reader_for(output, ConvertOptions(), settings)
    return output


def image_open(path: Path) -> Image.Image:
    try:
        with Image.open(path) as source:
            if source.width * source.height > 40_000_000:
                raise ConversionError('Слишком большое изображение: допустимо до 40 мегапикселей.')
            image = ImageOps.exif_transpose(source).copy()
            image.load()
            return image
    except (UnidentifiedImageError, OSError, Image.DecompressionBombError):
        raise ConversionError('Не удалось открыть изображение или превышен лимит его размеров.') from None


def rgb_on_white(image: Image.Image) -> Image.Image:
    rgba = image.convert('RGBA')
    base = Image.new('RGB', image.size, 'white')
    base.paste(rgba, mask=rgba.getchannel('A'))
    rgba.close()
    return base


def process_images(tool: str, inputs: list[Path], outdir: Path, options: ConvertOptions, progress: Progress) -> Path:
    if tool == 'images-to-pdf':
        output = outdir / 'result.pdf'
        document = canvas.Canvas(str(output))
        for i, path in enumerate(inputs):
            with image_open(path) as image:
                width, height = image.size
                if options.page_size == 'a4':
                    page_width, page_height = (841.89, 595.28) if width > height else (595.28, 841.89)
                    factor = min((page_width - 40) / width, (page_height - 40) / height)
                else:
                    factor = min(1, 2000 / max(width, height))
                    page_width, page_height = width * factor, height * factor
                document.setPageSize((page_width, page_height))
                target_width, target_height = width * factor, height * factor
                document.drawImage(ImageReader(image), (page_width - target_width) / 2, (page_height - target_height) / 2, target_width, target_height, mask='auto')
                document.showPage()
            progress(round(85 * (i + 1) / len(inputs)), f'Изображение {i + 1} из {len(inputs)}')
        document.save()
        return output
    with image_open(inputs[0]) as image:
        if tool == 'image-resize':
            image.thumbnail((options.width, options.height), Image.Resampling.LANCZOS)
        output = outdir / f'result.{options.format}'
        # New pixel-only image avoids carrying EXIF, comments and other metadata.
        mode = 'RGBA' if 'A' in image.getbands() or 'transparency' in image.info else 'RGB'
        pixels = image.convert(mode)
        if options.format == 'jpg':
            flattened = rgb_on_white(pixels)
            flattened.save(output, 'JPEG', quality=options.quality, optimize=True)
            flattened.close()
        elif options.format == 'png':
            clean = Image.new(mode, pixels.size)
            clean.paste(pixels)
            clean.save(output, 'PNG', optimize=True)
            clean.close()
        else:
            pixels.save(output, 'WEBP', quality=options.quality, method=4, exif=b'', icc_profile=b'')
        pixels.close()
    return output


def convert(tool: str, inputs: list[Path], outdir: Path, options: ConvertOptions, settings: Settings, progress: Progress) -> tuple[Path, list[str]]:
    outdir.mkdir(parents=True, exist_ok=True)
    warnings: list[str] = []
    for path in inputs:
        validate_container(path)
    progress(5, 'Проверяем документ')
    if tool in ('pptx-to-pdf', 'docx-to-pdf', 'excel-to-pdf'):
        return office_to_pdf(inputs[0], outdir, settings, progress), ['Верстка может отличаться, если исходные шрифты отсутствуют на сервере.']
    if tool in ('images-to-pdf', 'image-convert', 'image-compress', 'image-resize'):
        return process_images(tool, inputs, outdir, options, progress), warnings

    source = reader_for(inputs[0], options, settings)
    count = len(source.pages)
    output = outdir / 'result.pdf'
    indexes = parse_pages(options.pages, count)
    writer = PdfWriter()
    if tool == 'merge-pdf':
        total_pages = 0
        for i, path in enumerate(inputs):
            reader = reader_for(path, options, settings)
            total_pages += len(reader.pages)
            if total_pages > settings.max_pages:
                raise ConversionError(f'Итоговый документ не должен превышать {settings.max_pages} страниц.')
            writer.append(reader, import_outline=False)
            progress(round(85 * (i + 1) / len(inputs)), f'Документ {i + 1} из {len(inputs)}')
    elif tool == 'split-pdf':
        parts = []
        for i, page in enumerate(source.pages):
            part = PdfWriter()
            part.add_page(page)
            parts.append(write_pdf(part, outdir / f'page-{i+1:03d}.pdf'))
            progress(round(85 * (i + 1) / count), f'Страница {i + 1} из {count}')
        return make_zip(parts, outdir / 'pages.zip'), warnings
    elif tool in ('extract-pages', 'remove-pages'):
        selected = parse_pages(options.pages, count, allow_empty=False)
        if tool == 'remove-pages':
            selected = [i for i in range(count) if i not in set(selected)]
        if not selected:
            raise ConversionError('В документе должна остаться хотя бы одна страница.')
        for i in selected:
            writer.add_page(source.pages[i])
    elif tool == 'rotate-pdf':
        for i, page in enumerate(source.pages):
            if i in indexes:
                page.rotate(options.angle)
            writer.add_page(page)
    elif tool in ('compress-pdf', 'protect-pdf', 'unlock-pdf', 'pdf-metadata'):
        writer.clone_document_from_reader(source)
        if tool == 'compress-pdf':
            for i, page in enumerate(writer.pages):
                page.compress_content_streams()
                progress(round(85 * (i + 1) / count), f'Оптимизация страницы {i + 1} из {count}')
            writer.compress_identical_objects(remove_identicals=True, remove_orphans=True)
        elif tool == 'protect-pdf':
            if len(options.new_password) < 8:
                raise ConversionError('Задайте новый пароль длиной не менее 8 символов.')
            import secrets
            writer.encrypt(options.new_password, secrets.token_urlsafe(32), algorithm='AES-256')
        elif tool == 'pdf-metadata':
            from pypdf.generic import NameObject
            writer.add_metadata({})
            writer._info = None
            if '/Metadata' in writer.root_object:
                del writer.root_object[NameObject('/Metadata')]
            for page in writer.pages:
                if '/Metadata' in page:
                    del page[NameObject('/Metadata')]
            warnings.append('Удалены стандартные метаданные и XMP. Содержимое страниц и вложения не обезличиваются.')
    elif tool in ('watermark-pdf', 'number-pages'):
        if tool == 'watermark-pdf' and not options.text.strip():
            raise ConversionError('Введите текст водяного знака.')
        font = stamp_font()
        if font == 'Helvetica' and any(ord(char) > 127 for char in options.text):
            raise ConversionError('На сервере нет шрифта для кириллицы. Настройте STAMP_FONT_PATH.')
        from reportlab.pdfbase.pdfmetrics import stringWidth
        for i, page in enumerate(source.pages):
            page.transfer_rotation_to_content()
            width, height = float(page.mediabox.width), float(page.mediabox.height)
            overlay_bytes = io.BytesIO()
            overlay = canvas.Canvas(overlay_bytes, pagesize=(width, height))
            if tool == 'watermark-pdf':
                fontsize = min(48, width / 12)
                actual = stringWidth(options.text, font, fontsize)
                if actual > width * 0.85:
                    fontsize *= width * 0.85 / actual
                overlay.saveState()
                overlay.translate(width / 2, height / 2)
                overlay.rotate(30)
                overlay.setFillAlpha(0.22)
                overlay.setFillColorRGB(0.25, 0.27, 0.32)
                overlay.setFont(font, fontsize)
                overlay.drawCentredString(0, 0, options.text)
                overlay.restoreState()
            else:
                overlay.setFont(font, 10)
                overlay.setFillColorRGB(0.25, 0.27, 0.32)
                overlay.drawCentredString(width / 2, min(18, height / 10), str(options.start + i))
            overlay.save()
            overlay_bytes.seek(0)
            page.merge_page(PdfReader(overlay_bytes).pages[0])
            writer.add_page(page)
            progress(round(85 * (i + 1) / count), f'Страница {i + 1} из {count}')
    elif tool == 'pdf-to-text':
        texts = []
        for i, index in enumerate(indexes):
            texts.append(source.pages[index].extract_text() or '')
            progress(round(85 * (i + 1) / len(indexes)), f'Страница {i + 1} из {len(indexes)}')
        if not any(text.strip() for text in texts):
            raise ConversionError('Текстовый слой не найден. Для скана используйте «Распознать скан».')
        text_output = outdir / 'result.txt'
        text_output.write_text('\n\n\f\n\n'.join(texts), encoding='utf-8')
        return text_output, warnings
    elif tool in ('pdf-to-png', 'pdf-to-jpg'):
        extension = 'png' if tool == 'pdf-to-png' else 'jpg'
        files = []
        for i, (image, _) in enumerate(render_pages(inputs[0], indexes, options.dpi, options.password, progress)):
            target = outdir / f'page-{indexes[i]+1:03d}.{extension}'
            image.save(target, **({'quality': 92} if extension == 'jpg' else {}))
            image.close()
            files.append(target)
        return (files[0] if len(files) == 1 else make_zip(files, outdir / 'pages.zip')), warnings
    elif tool == 'pdf-to-pptx':
        from pptx import Presentation
        from pptx.util import Inches
        deck = Presentation()
        for i, (image, (page_width, page_height)) in enumerate(render_pages(inputs[0], list(range(count)), options.dpi, options.password, progress)):
            if i == 0:
                # PowerPoint has a single page size for the entire deck, max 56 inches.
                factor = min(1, 50 * 72 / max(page_width, page_height))
                if min(page_width, page_height) * factor < 72:
                    image.close()
                    raise ConversionError('Пропорции страницы выходят за пределы допустимых размеров PowerPoint.')
                deck.slide_width = Inches(page_width * factor / 72)
                deck.slide_height = Inches(page_height * factor / 72)
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            ratio = min(deck.slide_width / image.width, deck.slide_height / image.height)
            width, height = round(image.width * ratio), round(image.height * ratio)
            buffer = io.BytesIO()
            image.save(buffer, format='PNG')
            image.close()
            buffer.seek(0)
            slide.shapes.add_picture(buffer, (deck.slide_width - width) // 2, (deck.slide_height - height) // 2, width=width, height=height)
        deck.core_properties.title = 'Converted with IMONSOFT'
        deck.core_properties.subject = 'Image-based slides; original PDF elements are not separately editable.'
        target = outdir / 'result.pptx'
        deck.save(target)
        return target, ['Слайды содержат изображения страниц. Текст и объекты не редактируются отдельно.']
    elif tool == 'ocr-pdf':
        executable = shutil.which('tesseract')
        if not executable:
            raise ConversionError('На сервере не установлен Tesseract.')
        parts = []
        for i, (image, _) in enumerate(render_pages(inputs[0], list(range(count)), 144, options.password, progress)):
            raster = outdir / f'ocr-{i:03d}.png'
            image.save(raster, dpi=(144, 144))
            image.close()
            base = outdir / f'ocr-{i:03d}'
            try:
                process = subprocess.run([executable, str(raster), str(base), '-l', options.language, '--dpi', '144', 'pdf'], stdin=subprocess.DEVNULL, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=45, check=False)
            except subprocess.TimeoutExpired:
                raise ConversionError('Распознавание страницы превысило лимит времени.') from None
            raster.unlink(missing_ok=True)
            part = base.with_suffix('.pdf')
            if process.returncode != 0 or not part.exists():
                raise ConversionError('Распознавание не выполнено. Проверьте языковые пакеты Tesseract.')
            parts.append(part)
            writer.append(str(part), import_outline=False)
        write_pdf(writer, output)
        for part in parts:
            part.unlink(missing_ok=True)
        return output, ['Страницы растрированы. Текст распознан автоматически и может содержать ошибки.']
    else:
        raise ConversionError('Этот инструмент пока не поддерживается.')
    result = write_pdf(writer, output)
    if tool == 'compress-pdf' and result.stat().st_size >= inputs[0].stat().st_size:
        shutil.copyfile(inputs[0], result)
        warnings.append('Документ уже оптимизирован. Сохранён оригинал без увеличения размера.')
    return result, warnings
