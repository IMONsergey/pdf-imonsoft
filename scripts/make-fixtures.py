"""Create synthetic test documents only. Do not commit generated files."""
from pathlib import Path
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas

root = Path(__file__).resolve().parents[1] / '.data' / 'fixtures'
root.mkdir(parents=True, exist_ok=True)
c = canvas.Canvas(str(root / 'sample.pdf'), pagesize=(960, 540))
for index in range(1, 4):
    c.setFillColorRGB(.08, .12, .22)
    c.rect(0, 0, 960, 540, fill=1, stroke=0)
    c.setFillColorRGB(1, 1, 1)
    c.setFont('Helvetica', 36)
    c.drawString(70, 390, f'IMONSOFT / Page {index}')
    c.setFont('Helvetica', 18)
    c.drawString(70, 325, 'Synthetic conversion test. 960 x 540 pt.')
    c.showPage()
c.save()
ppt = Presentation()
ppt.slide_width = Inches(13.333333)
ppt.slide_height = Inches(7.5)
slide = ppt.slides.add_slide(ppt.slide_layouts[6])
slide.shapes.add_textbox(Inches(1), Inches(1), Inches(10), Inches(1)).text = 'IMONSOFT PowerPoint conversion test'
ppt.save(root / 'sample.pptx')
im = Image.new('RGB', (1200, 800), '#204fff')
ImageDraw.Draw(im).rectangle((150, 150, 1050, 650), fill='white')
im.save(root / 'sample.png')
print(root)

# Password and pagination fixtures are synthetic and contain no user credentials.
from pypdf import PdfReader, PdfWriter
writer = PdfWriter(clone_from=root / 'sample.pdf')
writer.encrypt('fixture-password-123', algorithm='AES-256')
writer.write(root / 'protected.pdf')
writer.close()
writer = PdfWriter()
for index in range(27):
    writer.add_page(PdfReader(root / 'sample.pdf').pages[index % 3])
writer.write(root / 'many-pages.pdf')
writer.close()

# A real multi-megabyte image PDF for end-to-end compression measurements.
import sys
import tempfile
import shutil
repository = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(repository / 'services/api'), str(repository / 'services/api/tests')]
from compression_fixtures import image_pdf
from app.compression import compress_pdf
from app.config import Settings
from app.models import ConvertOptions
image_pdf(root / 'compression-heavy.pdf', jpeg=True, nested=True)
with tempfile.TemporaryDirectory(prefix='imonsoft-fixture-') as temporary:
    first = Path(temporary) / 'first'; first.mkdir()
    second = Path(temporary) / 'second'; second.mkdir()
    optimized, _ = compress_pdf(root / 'compression-heavy.pdf', first, ConvertOptions(), Settings(), lambda *_: None)
    stable, _ = compress_pdf(optimized, second, ConvertOptions(), Settings(), lambda *_: None)
    shutil.copyfile(stable, root / 'compression-optimized.pdf')
