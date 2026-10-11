"""Exercise every catalog tool against the running local API using synthetic inputs.

Run from the repository with .venv/bin/python scripts/check-api-matrix.py.
Outputs contain only synthetic fixture metadata; jobs are deleted after checking.
"""
from __future__ import annotations
import io
import json
import tempfile
import time
import zipfile
from pathlib import Path
import httpx
from PIL import Image
from docx import Document
from pptx import Presentation
from pypdf import PdfReader, PdfWriter
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas
import xlsxwriter

ROOT = Path(__file__).resolve().parents[1]
BASE = 'http://127.0.0.1:8080'
HEADERS = {'X-IMONSOFT-Client': 'web'}
PASSWORD = 'synthetic-matrix-password'


def fixtures(directory: Path) -> dict[str, Path]:
    picture = directory / 'image.png'
    image = Image.new('RGBA', (180, 100), (30, 110, 180, 230))
    image.save(picture)
    pdf = directory / 'document.pdf'
    document = canvas.Canvas(str(pdf), pagesize=(720, 405))
    for index in range(3):
        document.setFont('Helvetica', 28)
        document.drawString(40, 330, f'IMONSOFT MATRIX PAGE {index + 1}')
        document.drawImage(ImageReader(image), 40, 90, 180, 100, mask='auto')
        document.showPage()
    document.save(); image.close()
    protected = directory / 'protected.pdf'
    writer = PdfWriter(clone_from=pdf)
    writer.encrypt(PASSWORD, algorithm='AES-256'); writer.write(protected); writer.close()
    docx = directory / 'document.docx'
    word = Document(); word.add_heading('IMONSOFT matrix', 0); word.add_paragraph('Synthetic Word document.'); word.save(docx)
    pptx = directory / 'slides.pptx'
    deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(914400, 914400, 7315200, 914400).text = 'IMONSOFT matrix presentation'
    deck.save(pptx)
    xlsx = directory / 'sheet.xlsx'
    workbook = xlsxwriter.Workbook(xlsx); sheet = workbook.add_worksheet('Matrix')
    sheet.write_row('A1', ['Item', 'Count']); sheet.write_row('A2', ['Synthetic', 3]); sheet.set_column('A:B', 20); workbook.close()
    return {'pdf': pdf, 'protected': protected, 'image': picture, 'word': docx, 'office': pptx, 'excel': xlsx}


def verify_output(extension: str, content: bytes, tool: str) -> None:
    assert content, f'{tool}: empty output'
    stream = io.BytesIO(content)
    if extension == '.pdf':
        reader = PdfReader(stream)
        if reader.is_encrypted:
            assert reader.decrypt(PASSWORD)
        assert len(reader.pages) > 0
    elif extension == '.pptx':
        assert len(Presentation(stream).slides) > 0
    elif extension == '.docx':
        assert any(paragraph.text.strip() for paragraph in Document(stream).paragraphs)
    elif extension == '.zip':
        with zipfile.ZipFile(stream) as archive:
            assert archive.namelist() and archive.testzip() is None
            assert all(not name.startswith('/') and '..' not in Path(name).parts for name in archive.namelist())
    elif extension in ('.png', '.jpg', '.jpeg', '.webp'):
        with Image.open(stream) as image:
            image.verify()
    elif extension == '.json':
        assert isinstance(json.loads(content), dict)
    else:
        assert content.decode('utf-8').strip()


def main() -> int:
    report = []
    catalog = json.loads((ROOT / 'shared/tools.json').read_text())
    output = ROOT / '.data' / 'qa' / 'v02-api-matrix.json'
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='imonsoft-matrix-') as temporary, httpx.Client(base_url=BASE, headers=HEADERS, timeout=30) as client:
        config = client.get('/api/v1/config'); config.raise_for_status()
        assert len(config.json()['tools']) == len(catalog)
        inputs = fixtures(Path(temporary))
        for definition in catalog:
            tool = definition['id']; started = time.monotonic(); job_id = None
            try:
                source = inputs['pdf']; selected = []
                options = {}
                if tool == 'pptx-to-pdf': source = inputs['office']
                elif tool == 'docx-to-pdf': source = inputs['word']
                elif tool == 'excel-to-pdf': source = inputs['excel']
                elif tool == 'unlock-pdf': source = inputs['protected']; options['password'] = PASSWORD
                elif tool.startswith('image'): source = inputs['image']
                if tool == 'protect-pdf': options['new_password'] = PASSWORD
                if tool == 'watermark-pdf': options['text'] = 'MATRIX REVIEW'
                if tool in ('extract-pages', 'remove-pages'): options['pages'] = '2'
                if tool == 'organize-pdf': options['page_plan'] = [{'page': 3, 'rotation': 90}, {'page': 1}]
                if tool == 'ocr-pdf': options['language'] = 'eng'
                selected = [source] * (2 if tool == 'merge-pdf' else 1)
                payload = [('files', (file.name, file.read_bytes())) for file in selected]
                response = client.post('/api/v1/jobs', data={'tool': tool, 'options': json.dumps(options)}, files=payload)
                response.raise_for_status(); job_id = response.json()['id']
                deadline = time.monotonic() + 180
                while time.monotonic() < deadline:
                    status = client.get(f'/api/v1/jobs/{job_id}'); status.raise_for_status(); job = status.json()
                    if job['state'] in ('done', 'error', 'cancelled'): break
                    time.sleep(0.2)
                else: raise AssertionError('API job deadline exceeded')
                assert job['state'] == 'done', job.get('error') or job['state']
                result = client.get(job['downloadUrl']); result.raise_for_status()
                verify_output(job['outputExtension'], result.content, tool)
                report.append({'tool': tool, 'status': 'passed', 'extension': job['outputExtension'], 'bytes': len(result.content), 'seconds': round(time.monotonic() - started, 2)})
                print(tool + ': passed', flush=True)
            except Exception as error:
                report.append({'tool': tool, 'status': 'failed', 'error': str(error)})
                print(tool + ': FAILED: ' + str(error), flush=True)
            finally:
                if job_id:
                    client.delete(f'/api/v1/jobs/{job_id}')
    data = {'base': BASE, 'tools': len(report), 'passed': sum(item['status'] == 'passed' for item in report), 'results': report}
    output.write_text(json.dumps(data, ensure_ascii=False, indent=2))
    print(f"{data['passed']}/{len(report)} tools passed. Report: {output}")
    return 0 if data['passed'] == len(catalog) else 1


if __name__ == '__main__':
    raise SystemExit(main())
